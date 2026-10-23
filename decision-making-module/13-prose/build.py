# -*- coding: utf-8 -*-
import re, sys
sys.path.insert(0,'.')
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from body import HEADER, P, CAPTIONS, REFS

SRC = "/root/.claude/uploads/1bea978d-276c-5fce-be57-d06d0e29354c/1a4f168f-MarshallWace_IRC_Final.docx"
def wc(s): return len([w for w in re.split(r'\s+', s.strip()) if w])
TOTAL = sum(wc(p) for p in P)

d = Document()
st = d.styles['Normal']; st.font.name = 'Verdana'; st.font.size = Pt(11)
pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE; pf.space_after = Pt(6)
for s in d.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)

def para(text="", size=11, bold=False, italic=False, align=None, after=6):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if align is not None: p.alignment = align
    r = p.add_run(text); r.font.name='Verdana'; r.font.size=Pt(size); r.bold=bold; r.italic=italic
    return p

for i, h in enumerate(HEADER):
    para(h, size=11, bold=(i == 2), after=0)
para("Word count: %d (the six paragraphs below; header, figure captions, references "
     "and appendix excluded)." % TOTAL, size=10, after=12)

FIGS = {3: '../11-figures/fig1-money.png', 6: '../11-figures/fig2-model.png'}
for n, text in enumerate(P, 1):
    para(text, after=(4 if n in FIGS else 10))
    if n in FIGS:
        d.add_picture(FIGS[n], width=Cm(15))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(CAPTIONS[n], size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

para("References", bold=True, after=6)
for r in REFS:
    p = para(r, size=10, after=4)
    p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1)

# carry Alex's appendix across verbatim -- it is his declaration, unaltered
d.add_page_break()
src = Document(SRC); sp = [x for x in src.paragraphs if x.text.strip()]
start = next(i for i, x in enumerate(sp) if x.text.startswith('Appendix:'))
for x in sp[start:]:
    t = x.text
    bold = t.startswith('Appendix:') or re.match(r'^(Prompt \d|Correction \d|Other corrections)', t)
    para(t, size=(12 if t.startswith('Appendix:') else 10), bold=bool(bold), after=4)

d.save('MarshallWace_IRC_prose.docx')
print("body word count :", TOTAL)
print("saved MarshallWace_IRC_prose.docx")
